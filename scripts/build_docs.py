"""
Build .docx files from the .md files in /docs.

Run from the project root:
    python scripts/build_docs.py
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS_DIR = Path(__file__).resolve().parent.parent / "docs"


def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)
    return h


def add_paragraph_with_formatting(doc, text):
    """Parse inline markdown (**bold**, *italic*, `code`, [text](url)) and add a paragraph."""
    p = doc.add_paragraph()
    # Split by formatting markers
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = p.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif part.startswith("[") and "](" in part:
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if match:
                run = p.add_run(match.group(1))
                run.font.color.rgb = RGBColor(0x1E, 0x58, 0xD6)
                run.font.underline = True
        else:
            p.add_run(part)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    return p


def add_numbered(doc, text):
    return doc.add_paragraph(text, style="List Number")


def add_quote(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.right_indent = Inches(0.5)
    for run in p.runs:
        run.italic = True
        run.font.color.rgb = RGBColor(0x5D, 0x6B, 0x80)
    return p


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    p.add_run("─" * 40).font.color.rgb = RGBColor(0xE3, 0xDD, 0xCF)
    return p


def parse_table(doc, lines):
    """Parse a markdown table starting at the given lines."""
    # lines[0] is header, lines[1] is separator, lines[2:] are rows
    header = [c.strip() for c in lines[0].strip("|").split("|")]
    rows = []
    for line in lines[2:]:
        row = [c.strip() for c in line.strip("|").split("|")]
        if len(row) == len(header):
            rows.append(row)
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = h
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    for r, row in enumerate(rows):
        for i, val in enumerate(row):
            table.rows[r + 1].cells[i].text = val
    doc.add_paragraph()


def add_code_block(doc, code):
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    return p


def md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()
    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    md = md_path.read_text(encoding="utf-8")
    lines = md.split("\n")

    i = 0
    in_code = False
    code_buffer = []

    while i < len(lines):
        line = lines[i]

        # Code block handling
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buffer = []
            else:
                in_code = False
                add_code_block(doc, "\n".join(code_buffer))
                code_buffer = []
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        stripped = line.strip()

        # Empty line
        if not stripped:
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,6})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            # Strip markdown formatting
            text = re.sub(r"\*+", "", text)
            add_heading(doc, text, level)
            i += 1
            continue

        # Horizontal rule
        if re.match(r"^---+$", stripped) or re.match(r"^\*\*\*+$", stripped):
            add_horizontal_rule(doc)
            i += 1
            continue

        # Tables
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|?\s*[-:|\s]+\|", lines[i + 1]):
            table_lines = []
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i])
                i += 1
            if table_lines:
                parse_table(doc, table_lines)
            continue

        # Blockquote
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            add_quote(doc, text)
            i += 1
            continue

        # Bulleted list
        m = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        if m:
            text = m.group(2)
            add_paragraph_with_formatting(doc, text)
            p = doc.paragraphs[-1]
            p.style = doc.styles["List Bullet"]
            i += 1
            continue

        # Numbered list
        m = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if m:
            text = m.group(2)
            add_paragraph_with_formatting(doc, text)
            p = doc.paragraphs[-1]
            p.style = doc.styles["List Number"]
            i += 1
            continue

        # Default: paragraph
        add_paragraph_with_formatting(doc, stripped)
        i += 1

    doc.save(str(docx_path))
    print(f"  Wrote {docx_path.name}")


def main():
    print("Building .docx files from .md sources...")
    if not DOCS_DIR.exists():
        print(f"  ERROR: docs/ folder not found at {DOCS_DIR}")
        return
    count = 0
    for md_file in sorted(DOCS_DIR.glob("*.md")):
        docx_file = md_file.with_suffix(".docx")
        try:
            md_to_docx(md_file, docx_file)
            count += 1
        except Exception as e:
            print(f"  ERROR with {md_file.name}: {e}")
    print(f"Done. {count} .docx files created in {DOCS_DIR}/")


if __name__ == "__main__":
    main()
